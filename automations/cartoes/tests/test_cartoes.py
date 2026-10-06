from contextlib import redirect_stderr, redirect_stdout
from datetime import date, datetime, timezone
from decimal import Decimal
from io import StringIO
import json
from pathlib import Path
import re
import runpy
import smtplib
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

import openpyxl

import cartoes as app


class FakeSMTP:
    def __init__(self, error_at=None, error=None):
        self.messages = []
        self.calls = 0
        self.error_at = error_at
        self.error = error
        self.closed = False

    def send_message(self, message):
        self.calls += 1
        if self.calls == self.error_at:
            raise self.error
        self.messages.append(message)
        return {}

    def close(self):
        self.closed = True


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.config = self.root / "config.json"
        self.config.write_text(json.dumps({
            "owner_name": "Titular de teste",
            "recipients": [
                {"sheet": "Ana", "participant": "Ana", "email": "ana@example.com"},
                {"sheet": "Dora", "participant": "Dora/Cae", "email": "dora@example.com"},
            ]
        }))
        self.settings = app.load_settings(self.config)
        self.path = self.root / "2026-09.xlsx"
        self.make_workbook()
        self.env = patch.dict(app.os.environ, {"GMAIL_USER": "owner@gmail.com", "GMAIL_APP_PASSWORD": "abcdefghijklmnop"})
        self.env.start()
        self.addCleanup(self.env.stop)

    def make_workbook(self):
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name in ("Black", "Latam"):
            s = book.create_sheet(name)
            s["A2"], s["F2"], s["G2"] = "PARTICIPANTES", "Ana", "Dora/Cae"
            s["A4"] = "CARTÃO AZUL"
            for r in (6, 7):
                s.cell(r, 1, date(2025, 1, 3))
                s.cell(r, 2, "Loja <script> & teste")
                s.cell(r, 3, "Parcela 10 de 12" if r == 6 else None)
                s.cell(r, 4, 20)
                s.cell(r, 6, 10)
                s.cell(r, 7, 10)
            s["A8"] = "CARTÃO BLACK"
            s["A9"], s["B9"], s["D9"], s["F9"] = date(2026, 9, 1), "Estorno", -2.35, -2.35
            s["A10"], s["D10"], s["F10"], s["G10"] = "TOTAL", "=SUM(D6:D9)", "=SUM(F6:F9)", "=SUM(G6:G9)"
            s["A12"], s["F12"], s["G12"] = "PAGAMENTOS", -5.65, -20
            s["A13"], s["F13"] = "PAGAMENTOS", -2
            s["A15"], s["F15"], s["G15"] = "TOTAL GERAL", "=SUM(F10:F14)", "=SUM(G10:G14)"
        book.create_sheet("Ana")
        s = book.create_sheet("Dora")
        s["A6"] = "#VALUE!"  # Cached summaries must not be needed.
        book.save(self.path)
        book.close()

    def change(self, fn):
        b = openpyxl.load_workbook(self.path)
        fn(b)
        b.save(self.path)
        b.close()

    def reports(self):
        return app.read_reports(self.path, self.settings, "black", "2026-09")

    def test_shared_purchases_credits_payments_and_old_installments(self):
        before = self.path.read_bytes()
        ana, dora = self.reports()
        self.assertEqual(len(ana.purchases), 3)
        self.assertEqual(ana.spending, Decimal("17.65"))
        self.assertEqual(ana.balance, Decimal("10.00"))
        self.assertEqual(ana.payments, (Decimal("-5.65"), Decimal("-2.00")))
        self.assertEqual(ana.purchases[0].date.year, 2025)
        self.assertEqual(ana.purchases[0].total, Decimal("20.00"))
        self.assertEqual(ana.purchases[1].installment, "—")
        self.assertTrue(dora.has_movement)
        self.assertEqual(dora.balance, Decimal("0.00"))
        self.assertEqual(before, self.path.read_bytes())

    def test_no_movement_is_not_an_error(self):
        self.change(lambda b: [setattr(b["Black"][f"G{r}"], "value", None) for r in (6, 7, 12)])
        self.assertFalse(self.reports()[1].has_movement)

    def test_payment_only_and_credit_balance(self):
        self.change(lambda b: [setattr(b["Black"][f"F{r}"], "value", None) for r in (6, 7, 9)])
        report = self.reports()[0]
        self.assertTrue(report.has_movement)
        self.assertEqual(report.balance, Decimal("-7.65"))
        self.assertIn("Crédito a seu favor", app.render_html(report))

    def test_no_payments_renders_zero_without_creating_movement(self):
        self.change(lambda b: [setattr(b["Black"][f"F{r}"], "value", None) for r in (12, 13)])
        report = self.reports()[0]
        self.assertEqual(report.payments, ())
        self.assertEqual(report.balance, report.spending)
        html = app.render_html(report)
        self.assertRegex(html, r"PAGAMENTOS</td><td[^>]*></td><td[^>]*>R\$ 0,00</td>")
        self.assertIn("PAGAMENTOS: R$ 0,00", app.render_text(report))
        self.change(lambda b: [setattr(b["Black"][f"F{r}"], "value", None) for r in (6, 7, 9)])
        self.assertFalse(self.reports()[0].has_movement)

    def test_numeric_rounding_and_bad_values(self):
        self.assertEqual(app._money(0.1 + 0.2, "F6"), Decimal("0.30"))
        self.assertEqual(app._money(1.005, "F6"), Decimal("1.01"))
        for value in ("10", "=SUM(A1:A3)", True, float("nan"), float("inf")):
            with self.subTest(value=str(value)), self.assertRaises(app.ReportError):
                app._money(value, "F6")

    def test_formula_in_source_is_rejected(self):
        self.change(lambda b: setattr(b["Black"]["F6"], "value", "=10"))
        with self.assertRaisesRegex(app.ReportError, "Black!F6"):
            self.reports()

    def test_invalid_date_and_purchase_total(self):
        self.change(lambda b: setattr(b["Black"]["A6"], "value", None))
        with self.assertRaisesRegex(app.ReportError, "data"):
            self.reports()
        self.make_workbook()
        self.change(lambda b: setattr(b["Black"]["D6"], "value", None))
        with self.assertRaisesRegex(app.ReportError, "Black!D6"):
            self.reports()

    def test_unknown_or_missing_sheets_and_headers_fail(self):
        self.change(lambda b: b.create_sheet("Novo amigo"))
        with self.assertRaisesRegex(app.ReportError, "sem cadastro"):
            self.reports()
        self.make_workbook()
        self.change(lambda b: b.remove(b["Dora"]))
        with self.assertRaisesRegex(app.ReportError, "ausentes"):
            self.reports()
        self.make_workbook()
        self.change(lambda b: setattr(b["Black"]["G2"], "value", "Outro"))
        with self.assertRaisesRegex(app.ReportError, "Dora"):
            self.reports()

    def test_duplicate_headers_and_bad_footer_fail(self):
        self.change(lambda b: setattr(b["Black"]["G2"], "value", "Ana"))
        with self.assertRaisesRegex(app.ReportError, "duplicado"):
            self.reports()
        self.make_workbook()
        self.change(lambda b: setattr(b["Black"]["A15"], "value", "Saldo"))
        with self.assertRaisesRegex(app.ReportError, "TOTAL GERAL"):
            self.reports()

    def test_header_columns_can_move(self):
        self.change(lambda b: b["Black"].insert_cols(6))
        self.assertEqual(self.reports()[0].balance, Decimal("10.00"))

    def test_html_escapes_content_and_keeps_five_columns(self):
        report = self.reports()[0]
        html = app.render_html(report)
        self.assertNotIn("<script>", html)
        self.assertIn("&lt;script&gt; &amp;", html)
        self.assertIn(app.BLUE, html)
        for label in ("Data", "Lançamento", "Parcelas", "Total", "Rateio", "TOTAL GERAL"):
            self.assertIn(label, html)
        self.assertIn("03/01/2025", html)
        self.assertIn("R$ 10,00", html)
        self.assertNotIn("2026-09.xlsx", html)
        self.assertNotIn("dora@example.com", html)

    def test_multipart_message_is_individual(self):
        report = self.reports()[0]
        msg = app.make_message(report, "owner@gmail.com", "Cartões", "ana@example.com")
        self.assertEqual(msg["To"], "ana@example.com")
        self.assertIsNone(msg["Cc"])
        self.assertIsNone(msg["Bcc"])
        self.assertEqual(msg.get_body(preferencelist=("plain",)).get_content_type(), "text/plain")
        self.assertEqual(msg.get_body(preferencelist=("html",)).get_content_type(), "text/html")
        for content in (msg.get_body(preferencelist=("plain",)).get_content(), msg.get_body(preferencelist=("html",)).get_content()):
            self.assertIn("Olá, Ana! Tudo bem?", content)
            self.assertIn("Sou o Hermes, assistente pessoal de Titular de teste.", content)
            self.assertIn("Segue abaixo o seu resumo, referente a 09/2026", content)
            self.assertIn("fale diretamente com Titular de teste.", content)

    def test_missing_owner_name_allows_preview_but_blocks_send(self):
        from dataclasses import replace
        report = replace(self.reports()[0], owner_name=None)
        self.assertIn("assistente pessoal de [seu nome]", app.render_html(report))
        with self.assertRaisesRegex(app.ReportError, "owner_name"):
            app.send_reports([report], self.settings, "hash", connector=lambda *_: self.fail("must not connect"))
        self.assertFalse(self.settings.state_dir.exists())

    def test_config_validation_and_relative_paths(self):
        self.assertEqual(self.settings.input_dir, self.root.resolve())
        data = json.loads(self.config.read_text())
        data["recipients"][1]["sheet"] = "Ana"
        self.config.write_text(json.dumps(data))
        with self.assertRaisesRegex(app.ReportError, "duplicado"):
            app.load_settings(self.config)
        self.assertFalse(app.valid_email("a@example.com\nBcc: b@example.com"))
        self.assertFalse(app.valid_email("a@example.com,b@example.com"))
        for email in ("a:b@example.com", "a@domain..com", ".a@example.com", "a@example.com\n", "a@-domain.com"):
            self.assertFalse(app.valid_email(email), email)
        self.assertTrue(app.valid_email("a.b+cartoes@example.com"))

    def test_scheduled_month_uses_brasilia_and_year_boundary(self):
        self.assertEqual(app.scheduled_month("black", datetime(2027, 1, 5, 12, tzinfo=timezone.utc)), "2026-12")
        self.assertEqual(app.scheduled_month("latam", datetime(2026, 10, 20, 12, tzinfo=timezone.utc)), "2026-09")
        with self.assertRaises(app.ReportError):
            app.scheduled_month("black", datetime(2026, 10, 5, 1, tzinfo=timezone.utc))
        for month in ("2026-13", "2026-9", "../2026-09", "0000-01"):
            with self.assertRaises(app.ReportError):
                app.validate_month(month)

    def test_dry_run_does_not_connect_or_write_ledger(self):
        with patch.object(app, "smtp_connect", side_effect=AssertionError("must not connect")), redirect_stdout(StringIO()):
            code = app.main(["--config", str(self.config), "--card", "black", "--month", "2026-09"])
        self.assertEqual(code, 0)
        self.assertEqual(len(list((self.root / "outputs/2026-09/black").glob("*.html"))), 3)
        self.assertFalse(self.settings.state_dir.exists())

    def test_preview_refresh_removes_only_previous_generated_files(self):
        folder = self.root / "preview"
        reports = self.reports()
        app.write_previews(reports, folder)
        custom = folder / "minha-nota.txt"
        custom.write_text("preservar")
        app.write_previews([reports[0]], folder)
        self.assertEqual(len(list(folder.glob("*.html"))), 2)  # Ana + index
        self.assertTrue(custom.exists())
        self.assertNotIn("Dora", (folder / "index.html").read_text())

    def test_gmail_connection_uses_validated_tls_and_closes_on_auth_failure(self):
        with patch.object(app.smtplib, "SMTP_SSL") as ssl_mock:
            connection = ssl_mock.return_value
            connection.login.side_effect = smtplib.SMTPAuthenticationError(535, b"denied")
            with self.assertRaises(smtplib.SMTPAuthenticationError):
                app.smtp_connect("owner@gmail.com", "password")
            args, kwargs = ssl_mock.call_args
            self.assertEqual(args, ("smtp.gmail.com", 465))
            self.assertTrue(kwargs["context"].check_hostname)
            self.assertEqual(kwargs["context"].verify_mode, app.ssl.CERT_REQUIRED)
            connection.close.assert_called_once()

    def test_hermes_launchers_use_config_and_scheduled_production_mode(self):
        root = Path(__file__).resolve().parents[1]
        for card in ("black", "latam"):
            default = Path("~/Automations/automations/cartoes/config.json").expanduser()
            for configured, expected in ((None, default), (str(self.config), self.config)):
                with self.subTest(card=card, configured=configured), patch.dict(app.os.environ), patch.object(app, "main", return_value=0) as main:
                    app.os.environ.pop("CARTOES_CONFIG", None)
                    if configured is not None:
                        app.os.environ["CARTOES_CONFIG"] = configured
                    with self.assertRaises(SystemExit) as exited:
                        runpy.run_path(str(root / f"hermes/cartao_{card}.py"), run_name="__main__")
                    self.assertEqual(exited.exception.code, 0)
                    main.assert_called_once_with(["--config", str(expected), "--card", card, "--scheduled", "--send"])

    def test_missing_file_notifies_only_owner_without_creating_ledger(self):
        self.path.unlink()
        smtp = FakeSMTP()
        with patch.object(app, "smtp_connect", return_value=smtp) as connect, redirect_stderr(StringIO()):
            self.assertEqual(app.main(["--config", str(self.config), "--card", "black", "--month", "2026-09", "--send"]), 1)
        connect.assert_called_once()
        self.assertEqual([m["To"] for m in smtp.messages], ["owner@gmail.com"])
        self.assertIn("[ERRO]", smtp.messages[0]["Subject"])
        self.assertIn("Nenhum resumo foi enviado", smtp.messages[0].get_content())
        self.assertFalse(self.settings.state_dir.exists())

    def test_corrupt_excel_at_open_or_lazy_read_notifies_only_owner(self):
        for failure in ("archive", "cell", "xml"):
            with self.subTest(failure=failure):
                self.make_workbook()
                if failure == "archive":
                    self.path.write_bytes(b"not an Excel archive")
                else:
                    with ZipFile(self.path) as source:
                        files = [(name, source.read(name)) for name in source.namelist()]
                    with ZipFile(self.path, "w") as target:
                        for name, data in files:
                            if name == "xl/worksheets/sheet1.xml":
                                if failure == "cell":
                                    data, count = re.subn(rb'(<c r="F6"[^>]*><v>)10(</v>)', rb'\g<1>PRIVATE_CELL_CONTENT\g<2>', data)
                                    self.assertEqual(count, 1)
                                else:
                                    data = data.replace(b"</sheetData>", b"</bad-sheetData>")
                            target.writestr(name, data)
                smtp, stderr = FakeSMTP(), StringIO()
                with patch.object(app, "smtp_connect", return_value=smtp) as connect, redirect_stderr(stderr):
                    code = app.main(["--config", str(self.config), "--card", "black", "--month", "2026-09", "--send"])
                self.assertEqual(code, 1)
                connect.assert_called_once()
                self.assertEqual(len(smtp.messages), 1)
                message = smtp.messages[0]
                self.assertEqual(message["To"], "owner@gmail.com")
                self.assertIsNone(message["Cc"])
                self.assertIsNone(message["Bcc"])
                self.assertEqual(list(message.iter_attachments()), [])
                self.assertIn("Excel", message.get_content())
                self.assertIn("Nenhum resumo foi enviado", message.get_content())
                for text in (stderr.getvalue(), message.as_string()):
                    self.assertNotIn("PRIVATE_CELL_CONTENT", text)
                    self.assertNotIn("Traceback", text)
                self.assertFalse(self.settings.state_dir.exists())
                self.assertTrue(smtp.closed)

    def test_invalid_values_or_recipients_notify_only_owner(self):
        for failure in ("value", "recipient"):
            with self.subTest(failure=failure):
                if failure == "value":
                    self.change(lambda b: setattr(b["Black"]["F6"], "value", "=10"))
                else:
                    self.make_workbook()
                    data = json.loads(self.config.read_text())
                    data["recipients"][1]["email"] = None
                    self.config.write_text(json.dumps(data))
                smtp = FakeSMTP()
                with patch.object(app, "smtp_connect", return_value=smtp), redirect_stderr(StringIO()):
                    code = app.main(["--config", str(self.config), "--card", "black", "--month", "2026-09", "--send"])
                self.assertEqual(code, 1)
                self.assertEqual([m["To"] for m in smtp.messages], ["owner@gmail.com"])
                self.assertFalse((self.settings.state_dir / "deliveries.sqlite3").exists())

    def test_invalid_config_uses_owner_credentials_from_adjacent_env(self):
        self.config.write_text("{invalid JSON")
        self.settings.secrets_file.write_text("GMAIL_USER=owner@gmail.com\nGMAIL_APP_PASSWORD=abcdefghijklmnop\n")
        self.settings.secrets_file.chmod(0o600)
        smtp = FakeSMTP()
        with patch.dict(app.os.environ, {}, clear=True), patch.object(app, "smtp_connect", return_value=smtp), redirect_stderr(StringIO()):
            code = app.main(["--config", str(self.config), "--card", "black", "--month", "2026-09", "--send"])
        self.assertEqual(code, 1)
        self.assertEqual([m["To"] for m in smtp.messages], ["owner@gmail.com"])
        self.assertIn("configuração", smtp.messages[0].get_content())

    def test_dry_run_errors_never_connect_to_gmail(self):
        self.path.write_bytes(b"not an Excel archive")
        for mode in ([], ["--dry-run"]):
            with self.subTest(mode=mode), patch.object(app, "smtp_connect") as connect, redirect_stderr(StringIO()):
                code = app.main(["--config", str(self.config), "--card", "black", "--month", "2026-09", *mode])
                self.assertEqual(code, 1)
                connect.assert_not_called()

    def test_alert_failure_does_not_retry_or_hide_original_error(self):
        self.path.unlink()
        for error in (smtplib.SMTPAuthenticationError(535, b"SECRET_SERVER_RESPONSE"), TimeoutError("SECRET_SERVER_RESPONSE")):
            with self.subTest(error=type(error).__name__):
                stderr = StringIO()
                with patch.object(app, "smtp_connect", side_effect=error) as connect, redirect_stderr(stderr):
                    code = app.main(["--config", str(self.config), "--card", "black", "--month", "2026-09", "--send"])
                self.assertEqual(code, 1)
                connect.assert_called_once()
                self.assertIn("ausente", stderr.getvalue())
                self.assertIn("Não foi possível confirmar o aviso", stderr.getvalue())
                self.assertNotIn("SECRET_SERVER_RESPONSE", stderr.getvalue())

    def test_alert_is_not_redirected_to_test_recipient(self):
        self.path.unlink()
        smtp = FakeSMTP()
        with patch.object(app, "smtp_connect", return_value=smtp), redirect_stderr(StringIO()):
            code = app.main(["--config", str(self.config), "--card", "black", "--month", "2026-09", "--send", "--test-to", "test@example.com"])
        self.assertEqual(code, 1)
        self.assertEqual([m["To"] for m in smtp.messages], ["owner@gmail.com"])

    def test_alert_rejection_does_not_retry(self):
        self.path.unlink()
        smtp = FakeSMTP(1, smtplib.SMTPRecipientsRefused({"owner@gmail.com": (550, b"rejected")}))
        stderr = StringIO()
        with patch.object(app, "smtp_connect", return_value=smtp) as connect, redirect_stderr(stderr):
            code = app.main(["--config", str(self.config), "--card", "black", "--month", "2026-09", "--send"])
        self.assertEqual(code, 1)
        connect.assert_called_once()
        self.assertEqual(smtp.calls, 1)
        self.assertTrue(smtp.closed)
        self.assertIn("Não foi possível confirmar o aviso", stderr.getvalue())

    def test_missing_gmail_credentials_preserves_failure_without_connecting(self):
        self.path.unlink()
        stderr = StringIO()
        with patch.dict(app.os.environ, {}, clear=True), patch.object(app, "smtp_connect") as connect, redirect_stderr(stderr):
            code = app.main(["--config", str(self.config), "--card", "black", "--month", "2026-09", "--send"])
        self.assertEqual(code, 1)
        connect.assert_not_called()
        self.assertIn("ausente", stderr.getvalue())
        self.assertIn("Não foi possível confirmar o aviso", stderr.getvalue())

    def test_unexpected_failure_alert_does_not_expose_exception_payload(self):
        smtp, stderr = FakeSMTP(), StringIO()
        with patch.object(app, "read_reports", side_effect=RuntimeError("PRIVATE_PAYLOAD")), patch.object(app, "smtp_connect", return_value=smtp), redirect_stderr(stderr):
            code = app.main(["--config", str(self.config), "--card", "black", "--month", "2026-09", "--send"])
        self.assertEqual(code, 1)
        self.assertEqual([m["To"] for m in smtp.messages], ["owner@gmail.com"])
        self.assertIn("RuntimeError", smtp.messages[0].get_content())
        self.assertNotIn("PRIVATE_PAYLOAD", stderr.getvalue() + smtp.messages[0].get_content())

    def test_partial_send_failure_sends_owner_alert_and_keeps_history(self):
        delivery = FakeSMTP(2, TimeoutError("PRIVATE_SERVER_RESPONSE"))
        alert = FakeSMTP()
        with patch.object(app, "smtp_connect", side_effect=[delivery, alert]), redirect_stderr(StringIO()):
            code = app.main(["--config", str(self.config), "--card", "black", "--month", "2026-09", "--send"])
        self.assertEqual(code, 1)
        self.assertEqual([m["To"] for m in delivery.messages], ["ana@example.com"])
        self.assertEqual([m["To"] for m in alert.messages], ["owner@gmail.com"])
        self.assertIn("Alguns resumos podem já ter sido aceitos", alert.messages[0].get_content())
        self.assertNotIn("Nenhum resumo foi enviado", alert.messages[0].get_content())
        self.assertNotIn("PRIVATE_SERVER_RESPONSE", alert.messages[0].get_content())
        with sqlite3.connect(self.settings.state_dir / "deliveries.sqlite3") as db:
            self.assertEqual(db.execute("SELECT sheet, status FROM deliveries ORDER BY sheet").fetchall(), [("Ana", "sent"), ("Dora", "unknown")])

    def test_history_shows_original_address_and_update_time(self):
        app.send_reports([self.reports()[0]], self.settings, "hash", connector=lambda *_: FakeSMTP())
        data = json.loads(self.config.read_text())
        data["recipients"][0]["email"] = "changed@example.com"
        self.config.write_text(json.dumps(data))
        stdout = StringIO()
        with redirect_stdout(stdout):
            code = app.main(["--config", str(self.config), "--history", "--month", "2026-09"])
        self.assertEqual(code, 0)
        self.assertIn("destinatário | atualização (UTC) | Message-ID", stdout.getvalue())
        self.assertIn("ana@example.com", stdout.getvalue())
        self.assertNotIn("changed@example.com", stdout.getvalue())
        self.assertRegex(stdout.getvalue(), r"\d{4}-\d{2}-\d{2}T.*\+00:00")

    def test_partial_failure_and_retry_skip_success(self):
        reports = self.reports()
        smtp = FakeSMTP(error_at=2, error=smtplib.SMTPDataError(550, b"Rejected"))
        with self.assertRaises(app.ReportError):
            app.send_reports(reports, self.settings, "hash", connector=lambda *_: smtp)
        self.assertTrue(smtp.closed)
        second = FakeSMTP()
        self.assertEqual(app.send_reports(reports, self.settings, "newhash", connector=lambda *_: second), (1, 1))
        self.assertEqual(second.messages[0]["To"], "dora@example.com")
        self.assertEqual(app.send_reports(reports, self.settings, "newhash", connector=lambda *_: self.fail("duplicate")), (0, 2))

    def test_timeout_is_unknown_and_not_retried(self):
        reports = self.reports()
        with self.assertRaisesRegex(app.ReportError, "incerto"):
            app.send_reports(reports, self.settings, "hash", connector=lambda *_: FakeSMTP(1, TimeoutError()))
        with self.assertRaisesRegex(app.ReportError, "incerto"):
            app.send_reports(reports, self.settings, "hash", connector=lambda *_: self.fail("must not connect"))
        with redirect_stdout(StringIO()):
            app.maintain_ledger(self.settings, "2026-09", "black", "Ana", "sent")
        smtp = FakeSMTP()
        self.assertEqual(app.send_reports(reports, self.settings, "hash", connector=lambda *_: smtp), (1, 1))

    def test_interruption_is_recovered_as_unknown(self):
        report = self.reports()[0]
        with app.run_lock(self.settings.state_dir):
            ledger = app.Ledger(self.settings.state_dir / "deliveries.sqlite3")
            ledger.claim(report, "<attempt@example.com>", "hash")
            ledger.close()
        with self.assertRaisesRegex(app.ReportError, "incerto"):
            app.send_reports([report], self.settings, "hash", connector=lambda *_: self.fail("must not connect"))

    def test_database_failure_after_smtp_acceptance_blocks_retry(self):
        report, smtp = self.reports()[0], FakeSMTP()
        with patch.object(app.Ledger, "finish", side_effect=sqlite3.OperationalError("disk error")):
            with self.assertRaises(sqlite3.OperationalError):
                app.send_reports([report], self.settings, "hash", connector=lambda *_: smtp)
        self.assertEqual(len(smtp.messages), 1)
        with self.assertRaisesRegex(app.ReportError, "incerto"):
            app.send_reports([report], self.settings, "hash", connector=lambda *_: self.fail("must not reconnect"))

    def test_test_delivery_redirects_without_production_history(self):
        smtp = FakeSMTP()
        app.send_reports(self.reports(), self.settings, "hash", test_to="owner@gmail.com", connector=lambda *_: smtp)
        self.assertEqual(len(smtp.messages), 2)
        for msg in smtp.messages:
            self.assertEqual(msg["To"], "owner@gmail.com")
            self.assertTrue(msg["Subject"].startswith("[TESTE]"))
        self.assertFalse((self.settings.state_dir / "deliveries.sqlite3").exists())

    def test_missing_recipient_blocks_whole_batch_before_smtp(self):
        report = self.reports()[1]
        bad = app.Report(report.card, report.month, app.Recipient("Dora", "Dora/Cae", None), report.purchases, report.payments)
        with self.assertRaisesRegex(app.ReportError, "Dora"):
            app.send_reports([self.reports()[0], bad], self.settings, "hash", connector=lambda *_: self.fail("must not connect"))
        self.assertFalse(self.settings.state_dir.exists())

    def test_lock_blocks_parallel_execution(self):
        with app.run_lock(self.settings.state_dir):
            with self.assertRaisesRegex(app.ReportError, "andamento"):
                with app.run_lock(self.settings.state_dir):
                    self.fail("second sender obtained lock")

    def test_lock_blocks_a_separate_process(self):
        script = """import sys
from pathlib import Path
import cartoes
try:
    with cartoes.run_lock(Path(sys.argv[1])):
        sys.exit(2)
except cartoes.ReportError:
    sys.exit(0)
"""
        with app.run_lock(self.settings.state_dir):
            result = subprocess.run([sys.executable, "-c", script, str(self.settings.state_dir)],
                                    cwd=Path(app.__file__).resolve().parent, capture_output=True, timeout=10)
        self.assertEqual(result.returncode, 0, result.stderr.decode())

    def test_smtp_auth_error_occurs_before_claim(self):
        with self.assertRaises(smtplib.SMTPAuthenticationError):
            app.send_reports(self.reports(), self.settings, "hash", connector=lambda *_: (_ for _ in ()).throw(smtplib.SMTPAuthenticationError(535, b"error")))
        with app.run_lock(self.settings.state_dir):
            ledger = app.Ledger(self.settings.state_dir / "deliveries.sqlite3")
            self.assertIsNone(ledger.get(self.reports()[0]))
            ledger.close()

    def test_secrets_file_permissions_and_spaced_app_password(self):
        self.settings.secrets_file.write_text("GMAIL_USER=owner@gmail.com\nGMAIL_APP_PASSWORD=abcd efgh ijkl mnop\n")
        self.settings.secrets_file.chmod(0o644)
        with self.assertRaisesRegex(app.ReportError, "600"):
            app.gmail_credentials(self.settings.secrets_file)
        self.settings.secrets_file.chmod(0o600)
        with patch.dict(app.os.environ, {}, clear=True):
            self.assertEqual(app.gmail_credentials(self.settings.secrets_file), ("owner@gmail.com", "abcdefghijklmnop"))


class AnonymousWorkbookTests(unittest.TestCase):
    def test_24_fictitious_blocks_without_private_workbook(self):
        fixture = json.loads((Path(__file__).parent / "fixtures/anonymous_reports.json").read_text(encoding="utf-8"))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = root / "config.json"
            config.write_text(json.dumps({"owner_name": "Proprietário fictício", "recipients": [
                {"sheet": p["sheet"], "participant": p["participant"], "email": None} for p in fixture["people"]
            ]}), encoding="utf-8")
            settings = app.load_settings(config)
            workbook = openpyxl.Workbook()
            workbook.remove(workbook.active)
            expected_purchases = {}
            for card, title in (("black", "Black"), ("latam", "Latam")):
                sheet = workbook.create_sheet(title)
                sheet["A2"] = "PARTICIPANTES"
                row = 4
                for column, person in enumerate(fixture["people"], 6):
                    sheet.cell(2, column, person["participant"])
                    expected_purchases[card, person["sheet"]] = []
                    sheet.cell(row, 1, "SEÇÃO DE COMPRAS")
                    row += 1
                    for value in person[card]["purchases"]:
                        amount = Decimal(value)
                        purchase = app.Purchase(date(2025, 1, 3), ("Descrição fictícia extensa com <símbolos> & detalhes " * 5).strip(),
                                                "2 de 12", amount * 2, amount)
                        sheet.cell(row, 1, purchase.date)
                        sheet.cell(row, 2, purchase.description)
                        sheet.cell(row, 3, purchase.installment)
                        sheet.cell(row, 4, float(purchase.total))
                        sheet.cell(row, column, float(amount))
                        expected_purchases[card, person["sheet"]].append(purchase)
                        row += 1
                sheet.cell(row, 1, "TOTAL")
                row += 1
                for column, person in enumerate(fixture["people"], 6):
                    for value in person[card]["payments"]:
                        sheet.cell(row, 1, "PAGAMENTOS")
                        sheet.cell(row, column, float(Decimal(value)))
                        row += 1
                sheet.cell(row, 1, "TOTAL GERAL")
            for person in fixture["people"]:
                workbook.create_sheet(person["sheet"])
            path = root / f"{fixture['month']}.xlsx"
            workbook.save(path)
            workbook.close()
            for card in ("black", "latam"):
                reports = app.read_reports(path, settings, card, fixture["month"])
                self.assertEqual(len(reports), 12)
                self.assertEqual(sum(r.has_movement for r in reports), fixture["expected_messages"][card])
                for report, person in zip(reports, fixture["people"], strict=True):
                    with self.subTest(card=card, sheet=person["sheet"]):
                        expected = person[card]
                        self.assertEqual(report.purchases, tuple(expected_purchases[card, person["sheet"]]))
                        self.assertEqual(report.payments, tuple(Decimal(v) for v in expected["payments"]))
                        self.assertEqual(report.spending, Decimal(expected["expected_spending"]))
                        self.assertEqual(report.balance, Decimal(expected["expected_balance"]))


class ProvidedWorkbookTests(unittest.TestCase):
    """Optional private fixture: never committed, included in local validation."""
    root = Path(__file__).resolve().parents[1]

    @unittest.skipUnless((root / "inputs/2026-09.xlsx").exists(), "Exemplo pessoal não está no checkout")
    def test_all_24_blocks_match_independent_cached_personal_summaries(self):
        settings = app.load_settings(self.root / "config.example.json")
        wb = openpyxl.load_workbook(self.root / "inputs/2026-09.xlsx", data_only=True)
        self.addCleanup(wb.close)
        for card, start, share_column, expected_count in (("black", 1, 5, 11), ("latam", 7, 11, 3)):
            reports = app.read_reports(self.root / "inputs/2026-09.xlsx", settings, card, "2026-09")
            self.assertEqual(len(reports), 12)
            self.assertEqual(sum(r.has_movement for r in reports), expected_count)
            for report in reports:
                with self.subTest(card=card, person=report.recipient.sheet):
                    self.assertLess(len(app.render_html(report).encode("utf-8")), 80_000)
                    source = wb[report.recipient.sheet]
                    purchases, payments = [], []
                    for row in range(6, source.max_row + 1):
                        when = source.cell(row, start).value
                        value = source.cell(row, share_column).value
                        if isinstance(when, datetime):
                            installment = source.cell(row, start + 2).value
                            purchases.append(app.Purchase(when.date(), source.cell(row, start + 1).value.strip(),
                                "—" if installment is None or installment == 0 else str(installment),
                                app._money(source.cell(row, start + 3).value, "fixture"), app._money(value, "fixture")))
                        elif when == "PAGAMENTOS":
                            payments.append(app._money(value, "fixture"))
                        elif when in ("TOTAL", "TOTAL GERAL"):
                            expected = Decimal(str(value or 0)).quantize(Decimal("0.01"), rounding="ROUND_HALF_UP")
                            self.assertEqual(report.spending if when == "TOTAL" else report.balance, expected)
                    self.assertEqual(report.purchases, tuple(purchases))
                    self.assertEqual(report.payments, tuple(payments))


if __name__ == "__main__":
    unittest.main()
