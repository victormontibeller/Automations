"""Package boundaries and installed entry points, with no Google access."""
import ast
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "src"


class PackageArchitectureTests(unittest.TestCase):
    def test_responsibilities_have_explicit_modules_and_no_cycles(self):
        required = {
            "models", "errors", "config", "workbook", "rendering", "previews",
            "scheduling", "locking", "ledger", "gmail", "drive", "delivery",
            "alerts", "service", "cli",
        }
        self.assertTrue(PACKAGE.is_dir(), "Application modules must live directly in src/")
        self.assertFalse((PACKAGE / "resumos_cartoes").exists(), "No extra source directory layer")
        self.assertFalse((PACKAGE / "cartoes.py").exists(), "Reuse the root CLI shim instead of duplicating it")
        self.assertTrue(required <= {p.stem for p in PACKAGE.glob("*.py")})
        graph = {}
        for path in PACKAGE.glob("*.py"):
            tree = ast.parse(path.read_text(encoding="utf-8"))
            dependencies = set()
            for node in ast.walk(tree):
                if isinstance(node, ast.ImportFrom) and node.level == 1:
                    dependencies.update([node.module.split(".")[0]] if node.module else [n.name for n in node.names])
                if isinstance(node, ast.Import):
                    self.assertFalse(any(n.name == "cartoes" for n in node.names), path.name)
                if isinstance(node, ast.Attribute):
                    self.assertNotEqual(ast.unparse(node), "sys.path", path.name)
            graph[path.stem] = dependencies
        def visit(module, ancestors):
            self.assertNotIn(module, ancestors, f"Circular imports: {ancestors} -> {module}")
            for dependency in graph.get(module, ()):
                visit(dependency, (*ancestors, module))
        for module in graph:
            visit(module, ())
        path = ROOT / "cartoes.py"
        tree = ast.parse(path.read_text(encoding="utf-8"))
        self.assertFalse(any(isinstance(n, (ast.FunctionDef, ast.ClassDef)) for n in ast.walk(tree)), path.name)
        self.assertNotIn("sys.path", path.read_text(encoding="utf-8"))

    def test_models_and_recalculation_import_without_google_or_cli(self):
        self.assertIsNotNone(importlib.util.find_spec("resumos_cartoes"), "Installable package missing")
        code = '''
import importlib.abc
import sys
class BlockExternal(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.split('.')[0] in {'google', 'googleapiclient', 'argparse'}:
            raise AssertionError('Unexpected dependency: ' + fullname)
sys.meta_path.insert(0, BlockExternal())
from decimal import Decimal
from resumos_cartoes.models import Recipient, Report
from resumos_cartoes.workbook import _money
report = Report('black', '2026-09', Recipient('Pessoa01', 'Participante01', None), (), (Decimal('-1.01'),))
assert report.balance == Decimal('-1.01')
assert _money(1.005, 'F6') == Decimal('1.01')
assert not any(name in sys.modules for name in ('resumos_cartoes.cli', 'resumos_cartoes.service', 'resumos_cartoes.gmail', 'resumos_cartoes.drive'))
'''
        result = subprocess.run([sys.executable, "-I", "-c", code], capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)


class InstalledCLITests(unittest.TestCase):
    def run_outside_checkout(self, args):
        with tempfile.TemporaryDirectory() as outside:
            env = os.environ.copy()
            env.pop("PYTHONPATH", None)
            return subprocess.run(args, cwd=outside, env=env, capture_output=True, text=True, timeout=30)

    def test_module_console_and_source_shim_help_outside_checkout(self):
        for command in (
            [sys.executable, "-I", "-m", "resumos_cartoes", "--help"],
            [str(Path(sys.executable).parent / "cartoes"), "--help"],
            [sys.executable, str(ROOT / "cartoes.py"), "--help"],
        ):
            with self.subTest(command=command):
                result = self.run_outside_checkout(command)
                self.assertEqual(result.returncode, 0, result.stderr)
                help_text = " ".join(result.stdout.split())
                self.assertIn("Black dia 5 = mês anterior", help_text)
                self.assertIn("Latam dia 20 = mês atual", help_text)
                for option in ("--send", "--scheduled", "--dry-run", "--history", "--resolve", "--test-to"):
                    self.assertIn(option, help_text)

    def test_installed_legacy_main_is_the_cli_main(self):
        code = '''
import cartoes
from resumos_cartoes.cli import main
assert cartoes.main is main
assert callable(main)
'''
        result = self.run_outside_checkout([sys.executable, "-I", "-c", code])
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_installed_packages_use_wheel_or_declared_editable_sources(self):
        code = '''
import importlib.metadata as metadata
import json
from pathlib import Path
import sys
import automation_core, resumos_cartoes
for name, module in (('automation-core', automation_core), ('resumos-cartoes', resumos_cartoes)):
    direct = json.loads(metadata.distribution(name).read_text('direct_url.json'))
    if not direct.get('dir_info', {}).get('editable'):
        assert Path(module.__file__).resolve().is_relative_to(Path(sys.prefix).resolve()), module.__file__
'''
        result = self.run_outside_checkout([sys.executable, "-I", "-c", code])
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_installed_preview_uses_synthetic_local_input(self):
        # Exercise the installed CLI end to end, not just argparse's help path.
        import openpyxl
        from datetime import date
        with tempfile.TemporaryDirectory() as outside:
            root = Path(outside)
            book = openpyxl.Workbook()
            book.remove(book.worksheets[0])
            for name in ("Black", "Latam"):
                sheet = book.create_sheet(name)
                sheet["F2"] = "Participante01"
                sheet["A4"], sheet["B4"], sheet["D4"], sheet["F4"] = date(2026, 9, 1), "Synthetic purchase", 2.01, 1.005
                sheet["A5"], sheet["A6"] = "TOTAL", "TOTAL GERAL"
            book.create_sheet("Pessoa01")
            book.save(root / "2026-09.xlsx")
            book.close()
            config = root / "config.json"
            config.write_text(json.dumps({
                "input_source": "local", "input_dir": ".",
                "google_token_file": "not-a-real-token.json",
                "recipients": [{"sheet": "Pessoa01", "participant": "Participante01", "email": None}],
            }), encoding="utf-8")
            result = subprocess.run(
                [sys.executable, "-I", "-m", "resumos_cartoes", "--config", str(config), "--card", "black", "--month", "2026-09"],
                cwd=root, capture_output=True, text=True, timeout=30,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("Nenhum e-mail enviado.", result.stdout)
            self.assertFalse((root / "var").exists())
            output = root / "outputs/2026-09/black"
            self.assertTrue((output / "index.html").is_file())
            text = next(output.glob("*.txt")).read_text(encoding="utf-8")
            self.assertIn("TOTAL GERAL: R$ 1,01", text)


if __name__ == "__main__":
    unittest.main()
