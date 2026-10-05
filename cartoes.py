"""Leitura determinística de rateios, prévias e envio individual via Gmail.

Uso: python cartoes.py --config config.example.json --card black --month 2026-09
"""

from __future__ import annotations

import argparse
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from email.message import EmailMessage
from email.utils import format_datetime, formataddr, make_msgid
import fcntl
import hashlib
from html import escape
import json
import os
from pathlib import Path
import re
import smtplib
import sqlite3
import ssl
import sys
from typing import BinaryIO, Callable
from zoneinfo import ZoneInfo

import openpyxl

ZONE = ZoneInfo("America/Sao_Paulo")
CARDS = {"black": ("Black", 5), "latam": ("Latam", 20)}
CENT = Decimal("0.01")
ZERO = Decimal("0.00")
BLUE = "#2F75B5"


class ReportError(Exception):
    """Erro operacional legível, sem credenciais ou conteúdo dos gastos."""


@dataclass(frozen=True)
class Recipient:
    sheet: str
    participant: str
    email: str | None
    enabled: bool = True


@dataclass(frozen=True)
class Settings:
    input_dir: Path
    output_dir: Path
    state_dir: Path
    secrets_file: Path
    sender_name: str
    recipients: tuple[Recipient, ...]
    owner_name: str | None = None


@dataclass(frozen=True)
class Purchase:
    date: date
    description: str
    installment: str
    total: Decimal
    share: Decimal


@dataclass(frozen=True)
class Report:
    card: str
    month: str
    recipient: Recipient
    purchases: tuple[Purchase, ...]
    payments: tuple[Decimal, ...]
    owner_name: str | None = None

    @property
    def spending(self) -> Decimal:
        return sum((p.share for p in self.purchases), ZERO)

    @property
    def balance(self) -> Decimal:
        return self.spending + sum(self.payments, ZERO)

    @property
    def has_movement(self) -> bool:
        return bool(self.purchases or self.payments)


def _clean_string(value: object, label: str) -> str:
    if not isinstance(value, str) or not value.strip() or any(c in value for c in "\r\n\x00"):
        raise ReportError(f"{label}: informe um texto válido em uma linha.")
    return value.strip()


def valid_email(value: str | None) -> bool:
    # One plain mailbox only: no display names, lists, BCC or header injection.
    if not isinstance(value, str) or len(value) > 254 or value.count("@") != 1:
        return False
    local, domain = value.split("@")
    if not re.fullmatch(r"[A-Za-z0-9!#$%&'*+/=?^_`{|}~.-]{1,64}", local):
        return False
    if local.startswith(".") or local.endswith(".") or ".." in local:
        return False
    labels = domain.split(".")
    return len(labels) >= 2 and all(re.fullmatch(r"[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?", label) for label in labels)


def load_settings(path: Path) -> Settings:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise ReportError("Não foi possível ler o arquivo de configuração JSON.") from exc
    if not isinstance(data, dict) or not isinstance(data.get("recipients"), list):
        raise ReportError("Configuração deve conter uma lista recipients.")
    recipients = []
    for item in data["recipients"]:
        if not isinstance(item, dict):
            raise ReportError("Cada destinatário deve ser um objeto JSON.")
        sheet = _clean_string(item.get("sheet"), "sheet")
        participant = _clean_string(item.get("participant"), "participant")
        email = item.get("email")
        if email is not None and not valid_email(email):
            raise ReportError(f"E-mail inválido no cadastro de {sheet}.")
        enabled = item.get("enabled", True)
        if not isinstance(enabled, bool):
            raise ReportError(f"enabled deve ser true ou false em {sheet}.")
        recipients.append(Recipient(sheet, participant, email, enabled))
    if not recipients:
        raise ReportError("Cadastre pelo menos uma aba pessoal.")
    for attr in ("sheet", "participant"):
        keys = [getattr(r, attr).casefold() for r in recipients]
        if len(keys) != len(set(keys)):
            raise ReportError(f"Cadastro duplicado no campo {attr}.")
    if any(r.sheet in ("Black", "Latam") for r in recipients):
        raise ReportError("As abas dos cartões não são destinatários.")
    base = path.resolve().parent

    def local_path(key: str, default: str) -> Path:
        value = _clean_string(data.get(key, default), key)
        return (base / Path(value).expanduser()).resolve()

    owner_name = data.get("owner_name")
    if owner_name is not None:
        owner_name = _clean_string(owner_name, "owner_name")
    return Settings(
        local_path("input_dir", "."), local_path("output_dir", "outputs"),
        local_path("state_dir", "var"), local_path("secrets_file", ".env"),
        _clean_string(data.get("sender_name", "Resumo dos cartões"), "sender_name"),
        tuple(recipients), owner_name,
    )


def validate_month(month: str) -> str:
    if not re.fullmatch(r"[0-9]{4}-[0-9]{2}", month):
        raise ReportError("Mês deve estar no formato AAAA-MM.")
    try:
        date.fromisoformat(month + "-01")
    except ValueError as exc:
        raise ReportError("Mês inválido.") from exc
    return month


def previous_month(today: date) -> str:
    return (today.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")


def scheduled_month(card: str, now: datetime | None = None) -> str:
    now = now or datetime.now(ZONE)
    local = now.astimezone(ZONE)
    if local.day != CARDS[card][1]:
        raise ReportError("Execução fora do dia previsto. Para recuperar um envio, informe --month AAAA-MM manualmente.")
    return previous_month(local.date())


def _money(value: object, cell: str) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ReportError(f"{cell}: valor monetário deve ser numérico, sem fórmula.")
    try:
        number = Decimal(str(value))
        if not number.is_finite():
            raise InvalidOperation
        return number.quantize(CENT, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError) as exc:
        raise ReportError(f"{cell}: valor monetário inválido.") from exc


def read_reports(path: Path | BinaryIO, settings: Settings, card: str, month: str) -> list[Report]:
    """Read primitive source cells; never evaluate formulas or trust their caches."""
    try:
        workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    except Exception as exc:
        raise ReportError(f"Não foi possível abrir {path.name}. Verifique se a cópia do Excel está completa.") from exc
    try:
        expected = {"Black", "Latam"} | {r.sheet for r in settings.recipients}
        unknown = set(workbook.sheetnames) - expected
        if unknown:
            raise ReportError("Abas sem cadastro: " + ", ".join(sorted(unknown)))
        missing = expected - set(workbook.sheetnames)
        if missing:
            raise ReportError("Abas ausentes: " + ", ".join(sorted(missing)))
        source = workbook[CARDS[card][0]]
        # Materialize once: read-only random access reparses XML repeatedly.
        rows = list(source.iter_rows())
        if len(rows) < 4 or len(rows[1]) < 6:
            raise ReportError(f"{source.title}: layout inválido; cabeçalhos esperados na linha 2 a partir de F.")
        headers: dict[str, int] = {}
        for i, cell in enumerate(rows[1][5:], start=5):
            if cell.value is None:
                continue
            key = str(cell.value).strip().casefold()
            if key in headers:
                raise ReportError(f"Cabeçalho duplicado em {source.title}, linha 2.")
            headers[key] = i
        labels = [(str(row[0].value or "").strip().upper(), i) for i, row in enumerate(rows)]
        totals = [i for label, i in labels if label == "TOTAL"]
        balances = [i for label, i in labels if label == "TOTAL GERAL"]
        if len(totals) != 1 or len(balances) != 1 or totals[0] >= balances[0]:
            raise ReportError(f"{source.title}: esperava uma linha TOTAL seguida de TOTAL GERAL.")
        reports = []
        for recipient in settings.recipients:
            if not recipient.enabled:
                continue
            column = headers.get(recipient.participant.casefold())
            if column is None:
                raise ReportError(f"Coluna de rateio não encontrada para {recipient.sheet} em {source.title}.")
            purchases, payments = [], []
            for i, row in enumerate(rows[3:], start=3):
                label = str(row[0].value or "").strip().upper()
                if label in {"TOTAL", "TOTAL GERAL"}:
                    continue
                value = row[column].value
                if value is None:
                    continue
                cell_label = f"{source.title}!{openpyxl.utils.get_column_letter(column + 1)}{i + 1}"
                share = _money(value, cell_label)
                if share == ZERO:
                    continue
                if label == "PAGAMENTOS":
                    if not totals[0] < i < balances[0]:
                        raise ReportError(f"{cell_label}: pagamento fora do bloco esperado.")
                    payments.append(share)
                    continue
                when = row[0].value
                if not isinstance(when, date) or i >= totals[0]:
                    raise ReportError(f"{cell_label}: rateio sem data de compra válida ou fora do bloco de compras.")
                description = row[1].value
                if not isinstance(description, str) or not description.strip() or row[1].data_type in {"f", "e"}:
                    raise ReportError(f"{source.title}!B{i + 1}: lançamento inválido.")
                installment = row[2].value
                if row[2].data_type in {"f", "e"}:
                    raise ReportError(f"{source.title}!C{i + 1}: parcela contém fórmula ou erro.")
                purchases.append(Purchase(
                    when.date() if isinstance(when, datetime) else when,
                    description.strip(), "—" if installment is None or installment == 0 else str(installment),
                    _money(row[3].value, f"{source.title}!D{i + 1}"), share,
                ))
            reports.append(Report(card, month, recipient, tuple(purchases), tuple(payments), settings.owner_name))
        return reports
    finally:
        workbook.close()


def brl(value: Decimal) -> str:
    formatted = f"{abs(value):,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    return ("-" if value < ZERO else "") + "R$ " + formatted


def subject(report: Report) -> str:
    return f"Cartão {report.card.upper()} {report.month} — {report.recipient.sheet}"


def greeting_paragraphs(report: Report) -> tuple[str, str, str, str]:
    owner = report.owner_name or "[seu nome]"
    year, month = report.month.split("-")
    return (
        f"Olá, {report.recipient.sheet}! Tudo bem?",
        f"Sou o Hermes, assistente pessoal de {owner}. Vou ajudar no envio dos resumos mensais dos cartões compartilhados.",
        f"Segue abaixo o seu resumo, referente a {month}/{year}, com os gastos, pagamentos registrados e saldo atualizado.",
        f"Se tiver alguma dúvida sobre os lançamentos, fale diretamente com {owner}.",
    )


def render_html(report: Report) -> str:
    """Use inline table styles supported by Gmail, without scripts or remote assets."""
    number = "text-align:right;white-space:nowrap;"
    cell_style = "padding:7px 6px;"
    row_open = '<tr style="border-bottom:1px solid #e5e7eb;vertical-align:top;">'

    def cell(text: str, style: str = "", colspan: int = 1) -> str:
        span = f' colspan="{colspan}"' if colspan > 1 else ""
        return f'<td{span} style="{cell_style}{style}">{escape(text)}</td>'

    rows = []
    for p in report.purchases:
        rows.append(row_open + cell(p.date.strftime("%d/%m/%Y"), "white-space:nowrap;")
                    + cell(p.description, "overflow-wrap:anywhere;word-break:break-word;")
                    + cell(p.installment) + cell(brl(p.total), number)
                    + cell(brl(p.share), number) + "</tr>")

    def summary_row(label: str, amount: Decimal, strong: bool = False) -> str:
        style = "font-weight:bold;" if strong else ""
        if strong:
            style += "border-top:2px solid #9ca3af;"
        # Leave the original purchase Total column empty in summary rows.
        return row_open + cell(label, style, 3) + cell("", style) + cell(brl(amount), number + style) + "</tr>"

    rows.append(summary_row("TOTAL", report.spending, True))
    rows.extend(summary_row("PAGAMENTOS", amount) for amount in report.payments)
    rows.append(summary_row("TOTAL GERAL", report.balance, True))
    status = "Saldo quitado." if report.balance == ZERO else "Crédito a seu favor." if report.balance < ZERO else "Saldo restante a pagar."
    title = f"Cartão {report.card.upper()} {report.month}"
    heads = "".join(
        f'<th scope="col" style="padding:10px 6px;font-size:14px;text-align:{alignment};border-bottom:1px solid #d1d5db;">{label}</th>'
        for label, alignment in (("Data", "left"), ("Lançamento", "left"), ("Parcelas", "left"), ("Total", "right"), ("Rateio", "right"))
    )
    greeting = greeting_paragraphs(report)
    introduction = "".join(
        f'<p style="font-size:16px;line-height:1.5;margin:8px 0 14px;">{escape(paragraph)}</p>'
        for paragraph in greeting[:3]
    )
    return f'''<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{escape(subject(report))}</title></head>
<body style="margin:0;padding:12px;background:#ffffff;color:#111111;font-family:Calibri,Arial,sans-serif;">
<table role="presentation" width="100%" style="width:100%;max-width:920px;margin:0 auto;border-collapse:collapse;"><tr><td>
{introduction}
<table width="100%" aria-label="Gastos do cartão" style="width:100%;border-collapse:collapse;font-family:Calibri,Arial,sans-serif;font-size:14px;">
<thead><tr><th colspan="5" bgcolor="{BLUE}" style="background:{BLUE};color:#111111;text-align:center;padding:12px;font-size:24px;font-weight:bold;">{escape(title)}</th></tr><tr>{heads}</tr></thead>
<tbody>{''.join(rows)}</tbody></table>
<p style="font-size:14px;margin:12px 0 4px;">{status} O valor individual está na coluna Rateio.</p>
<p style="font-size:16px;line-height:1.5;margin:14px 0 8px;">{escape(greeting[3])}</p>
</td></tr></table></body></html>'''


def render_text(report: Report) -> str:
    greeting = greeting_paragraphs(report)
    lines = ["\n\n".join(greeting[:3]), "", subject(report), "", "Data | Lançamento | Parcelas | Total | Rateio"]
    lines += [f"{p.date:%d/%m/%Y} | {p.description} | {p.installment} | {brl(p.total)} | {brl(p.share)}" for p in report.purchases]
    lines += ["", "TOTAL: " + brl(report.spending)]
    lines += ["PAGAMENTOS: " + brl(p) for p in report.payments]
    lines += ["TOTAL GERAL: " + brl(report.balance), "O valor individual está na coluna Rateio."]
    lines += ["", greeting[3]]
    return "\n".join(lines) + "\n"


def make_message(report: Report, sender: str, sender_name: str, target: str, *, test: bool = False, message_id: str | None = None) -> EmailMessage:
    message = EmailMessage()
    message["From"] = formataddr((sender_name, sender))
    message["To"] = target
    message["Subject"] = ("[TESTE] " if test else "") + subject(report)
    message["Date"] = format_datetime(datetime.now(timezone.utc))
    message["Message-ID"] = message_id or make_msgid(domain=sender.split("@")[-1])
    message.set_content(render_text(report))
    message.add_alternative(render_html(report), subtype="html")
    return message


def write_previews(reports: list[Report], output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True, mode=0o700)
    manifest = output / "manifest.json"
    if manifest.exists():
        try:
            previous = json.loads(manifest.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            raise ReportError("Manifesto das prévias inválido; use outra output_dir ou remova o manifesto.") from exc
        if not isinstance(previous, list):
            raise ReportError("Manifesto das prévias deve ser uma lista.")
        for filename in previous:
            # Only remove files previously generated by this script in this folder.
            if not isinstance(filename, str) or not re.fullmatch(r"[a-zA-Z0-9_-]+-[a-f0-9]{8}\.(html|txt)", filename):
                raise ReportError("Nome inválido no manifesto das prévias.")
            (output / filename).unlink(missing_ok=True)
    files, links = [], []
    for report in reports:
        slug = re.sub(r"[^a-zA-Z0-9_-]", "_", report.recipient.sheet)
        suffix = hashlib.sha256(report.recipient.sheet.encode()).hexdigest()[:8]
        stem = f"{slug}-{suffix}"
        for extension, content in (("html", render_html(report)), ("txt", render_text(report))):
            filename = f"{stem}.{extension}"
            (output / filename).write_text(content, encoding="utf-8")
            files.append(filename)
        links.append(f'<li><a href="{stem}.html">{escape(report.recipient.sheet)}</a> — saldo {escape(brl(report.balance))}</li>')
    (output / "index.html").write_text(
        '<!doctype html><html lang="pt-BR"><meta charset="utf-8"><title>Prévias dos cartões</title>'
        '<body style="font-family:Calibri,Arial,sans-serif;padding:20px"><h1>Prévias dos resumos</h1>'
        '<p>Nenhum e-mail foi enviado.</p><ul>' + ''.join(links) + '</ul></body></html>', encoding="utf-8")
    manifest.write_text(json.dumps(files, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def gmail_credentials(settings: Settings) -> tuple[str, str]:
    # A deliberately small KEY=VALUE parser. Never execute/source the file.
    values = {}
    if settings.secrets_file.exists():
        if settings.secrets_file.stat().st_mode & 0o077:
            raise ReportError("O arquivo de credenciais deve ter permissão 600 (chmod 600 .env).")
        for line in settings.secrets_file.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            key, sep, value = line.partition("=")
            if not sep or key.strip() not in {"GMAIL_USER", "GMAIL_APP_PASSWORD"}:
                raise ReportError("Arquivo de credenciais inválido; use apenas GMAIL_USER e GMAIL_APP_PASSWORD.")
            values[key.strip()] = value.strip().strip('"\'')
    for key in ("GMAIL_USER", "GMAIL_APP_PASSWORD"):
        if key in os.environ:
            values[key] = os.environ[key]
    user = values.get("GMAIL_USER", "")
    password = values.get("GMAIL_APP_PASSWORD", "").replace(" ", "")
    if not valid_email(user) or not re.fullmatch(r"[A-Za-z0-9]{16}", password):
        raise ReportError("Configure GMAIL_USER e uma senha de app Gmail de 16 caracteres no arquivo de credenciais ou no ambiente.")
    return user, password


@contextmanager
def run_lock(state_dir: Path):
    state_dir.mkdir(parents=True, exist_ok=True, mode=0o700)
    with (state_dir / "send.lock").open("a") as handle:
        try:
            fcntl.flock(handle, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as exc:
            raise ReportError("Já existe uma execução de envio ou manutenção do histórico em andamento.") from exc
        try:
            yield
        finally:
            fcntl.flock(handle, fcntl.LOCK_UN)


class Ledger:
    def __init__(self, path: Path):
        self.db = sqlite3.connect(path)
        self.db.row_factory = sqlite3.Row
        self.db.execute('''CREATE TABLE IF NOT EXISTS deliveries (
            month TEXT NOT NULL, card TEXT NOT NULL, sheet TEXT NOT NULL,
            status TEXT NOT NULL CHECK(status IN ('sending','sent','failed','unknown')),
            recipient TEXT NOT NULL, message_id TEXT NOT NULL, source_hash TEXT NOT NULL,
            updated_at TEXT NOT NULL, note TEXT NOT NULL DEFAULT '',
            PRIMARY KEY (month,card,sheet))''')
        # Caller holds the process lock: no live sender can own these entries.
        self.db.execute("UPDATE deliveries SET status='unknown', note='Execução interrompida; conferir no Gmail.' WHERE status='sending'")
        self.db.commit()

    def close(self):
        self.db.close()

    def get(self, report: Report):
        return self.db.execute("SELECT * FROM deliveries WHERE month=? AND card=? AND sheet=?", (report.month, report.card, report.recipient.sheet)).fetchone()

    def claim(self, report: Report, message_id: str, source_hash: str):
        self.db.execute('''INSERT INTO deliveries VALUES (?,?,?,'sending',?,?,?,?, '')
            ON CONFLICT(month,card,sheet) DO UPDATE SET status='sending',recipient=excluded.recipient,
            message_id=excluded.message_id,source_hash=excluded.source_hash,updated_at=excluded.updated_at,note='' ''',
            (report.month, report.card, report.recipient.sheet, report.recipient.email, message_id, source_hash, datetime.now(timezone.utc).isoformat()))
        self.db.commit()

    def finish(self, report: Report, status: str, note: str = ""):
        self.db.execute("UPDATE deliveries SET status=?,note=?,updated_at=? WHERE month=? AND card=? AND sheet=?",
                        (status, note, datetime.now(timezone.utc).isoformat(), report.month, report.card, report.recipient.sheet))
        self.db.commit()


def smtp_connect(user: str, password: str):
    connection = smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=30, context=ssl.create_default_context())
    try:
        connection.login(user, password)
    except BaseException:
        connection.close()
        raise
    return connection


def send_reports(reports: list[Report], settings: Settings, source_hash: str, *, test_to: str | None = None, connector: Callable | None = None) -> tuple[int, int]:
    if test_to is not None and not valid_email(test_to):
        raise ReportError("Endereço de teste inválido.")
    missing = [r.recipient.sheet for r in reports if not valid_email(r.recipient.email)]
    if missing and test_to is None:
        raise ReportError("Preencha os e-mails antes de enviar: " + ", ".join(missing))
    if not reports:
        return 0, 0
    if any(not r.owner_name or r.owner_name == "[seu nome]" for r in reports):
        raise ReportError("Preencha owner_name na configuração antes de enviar: é o nome apresentado na saudação do Hermes.")
    user, password = gmail_credentials(settings)
    with run_lock(settings.state_dir):
        ledger = None if test_to else Ledger(settings.state_dir / "deliveries.sqlite3")
        connection = None
        try:
            pending, skipped = [], 0
            for report in reports:
                prior = ledger.get(report) if ledger else None
                if prior and prior["status"] == "unknown":
                    raise ReportError(f"Envio de {report.recipient.sheet} tem resultado incerto. Confira o histórico e resolva antes de continuar.")
                if prior and prior["status"] == "sent":
                    skipped += 1
                else:
                    pending.append(report)
            if not pending:
                return 0, skipped
            connection = (connector or smtp_connect)(user, password)
            sent = 0
            for report in pending:
                message_id = make_msgid(domain=user.split("@")[-1])
                message = make_message(report, user, settings.sender_name, test_to or report.recipient.email,
                                       test=test_to is not None, message_id=message_id)
                if ledger:
                    ledger.claim(report, message_id, source_hash)
                try:
                    refused = connection.send_message(message)
                    if refused:
                        raise smtplib.SMTPRecipientsRefused(refused)
                except (smtplib.SMTPRecipientsRefused, smtplib.SMTPSenderRefused, smtplib.SMTPDataError) as exc:
                    if ledger:
                        ledger.finish(report, "failed", type(exc).__name__)
                    raise ReportError(f"SMTP recusou o envio de {report.recipient.sheet}. Mensagens anteriores permanecem registradas.") from exc
                except Exception as exc:
                    if ledger:
                        ledger.finish(report, "unknown", type(exc).__name__)
                    raise ReportError(f"Resultado incerto ao enviar para {report.recipient.sheet}; confira o Gmail antes de repetir.") from exc
                if ledger:
                    ledger.finish(report, "sent")
                sent += 1
            return sent, skipped
        finally:
            if connection is not None:
                connection.close()  # QUIT failures must not invalidate accepted mail.
            if ledger:
                ledger.close()


def maintain_ledger(settings: Settings, month: str | None, card: str | None, recipient: str | None, resolve: str | None) -> None:
    if not (settings.state_dir / "deliveries.sqlite3").exists():
        raise ReportError("Ainda não existe histórico de envios.")
    with run_lock(settings.state_dir):
        ledger = Ledger(settings.state_dir / "deliveries.sqlite3")
        try:
            if resolve:
                if not all((month, card, recipient)):
                    raise ReportError("--resolve exige --month, --card e --recipient.")
                rows = ledger.db.execute("SELECT * FROM deliveries WHERE month=? AND card=? AND sheet=?", (month, card, recipient)).fetchall()
                if len(rows) != 1 or rows[0]["status"] != "unknown":
                    raise ReportError("Só é possível resolver um envio com status unknown.")
                new_status = "sent" if resolve == "sent" else "failed"
                ledger.db.execute("UPDATE deliveries SET status=?,note=?,updated_at=? WHERE month=? AND card=? AND sheet=?",
                                  (new_status, "Conferido manualmente pelo proprietário.", datetime.now(timezone.utc).isoformat(), month, card, recipient))
                ledger.db.commit()
            rows = ledger.db.execute("SELECT month,card,sheet,status,message_id FROM deliveries ORDER BY month,card,sheet").fetchall()
            for row in rows:
                if (month is None or row["month"] == month) and (card is None or row["card"] == card) and (recipient is None or row["sheet"] == recipient):
                    print(" | ".join(str(value) for value in row))
        finally:
            ledger.close()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Resumos individuais dos cartões. Sem --send, apenas gera prévias.")
    parser.add_argument("--config", type=Path, default=Path("config.json"))
    parser.add_argument("--card", choices=CARDS)
    parser.add_argument("--month", help="AAAA-MM; obrigatório em execução manual")
    parser.add_argument("--recipient", help="Restringir a uma aba pessoal, pelo nome exato")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--dry-run", action="store_true", help="Gerar prévias, sem SMTP (padrão)")
    mode.add_argument("--send", action="store_true", help="Enviar mensagens")
    mode.add_argument("--history", action="store_true", help="Consultar o histórico")
    mode.add_argument("--resolve", choices=("sent", "not-sent"), help="Resolver resultado incerto após conferir o Gmail")
    parser.add_argument("--test-to", help="Com --send, redirecionar todos os resumos a um endereço de teste")
    parser.add_argument("--scheduled", action="store_true", help="Validar dia previsto e selecionar mês anterior")
    args = parser.parse_args(argv)
    try:
        settings = load_settings(args.config)
        if args.month:
            validate_month(args.month)
        if args.history or args.resolve:
            if args.scheduled or args.test_to:
                raise ReportError("Histórico não aceita --scheduled ou --test-to.")
            maintain_ledger(settings, args.month, args.card, args.recipient, args.resolve)
            return 0
        if not args.card:
            raise ReportError("Informe --card black ou --card latam.")
        if args.scheduled:
            if args.month:
                raise ReportError("--scheduled não aceita --month; recuperação deve ser manual.")
            month = scheduled_month(args.card)
        else:
            if not args.month:
                raise ReportError("Informe --month AAAA-MM para executar manualmente.")
            month = args.month
        if args.test_to and not args.send:
            raise ReportError("--test-to exige --send.")
        # Snapshot the bytes once so the audit hash and report share the same input.
        path = settings.input_dir / f"{month}.xlsx"
        try:
            from io import BytesIO
            source_bytes = path.read_bytes()
        except OSError as exc:
            raise ReportError(f"Arquivo {path.name} ausente ou inacessível na pasta configurada.") from exc
        # read_reports accepts a stream as well as a path; preserve the name for diagnostics.
        stream = BytesIO(source_bytes)
        stream.name = path.name
        reports = read_reports(stream, settings, args.card, month)
        if args.recipient:
            reports = [r for r in reports if r.recipient.sheet == args.recipient]
            if not reports:
                raise ReportError("Aba pessoal não cadastrada ou desativada.")
        active = [r for r in reports if r.has_movement]
        if args.send:
            count, skipped = send_reports(active, settings, hashlib.sha256(source_bytes).hexdigest(), test_to=args.test_to)
            print(f"{args.card.upper()} {month}: {count} mensagens {'de teste ' if args.test_to else ''}aceitas pelo SMTP; {skipped} já enviadas; {len(reports) - len(active)} sem movimento.")
        else:
            folder = settings.output_dir / month / args.card
            write_previews(active, folder)
            print(f"Prévia {args.card.upper()} {month}: {len(active)} resumos em {folder}; {len(reports) - len(active)} sem movimento. Nenhum e-mail enviado.")
        return 0
    except ReportError as exc:
        print(f"Erro: {exc}", file=sys.stderr)
        return 1
    except (smtplib.SMTPException, OSError, sqlite3.Error) as exc:
        # Exception payloads may contain SMTP addresses or credential server output.
        print(f"Erro operacional ({type(exc).__name__}). Verifique Gmail, permissões e acesso ao histórico.", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
