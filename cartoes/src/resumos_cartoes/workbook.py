"""Read and validate primitive Excel cells without trusting cached formulas."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import BinaryIO

import openpyxl

from .config import Settings
from .errors import ReportError
from .models import CENT, ZERO, Purchase, Report
from .scheduling import CARDS


def _money(value: object, cell: str) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ReportError(f"{cell}: valor monetário deve ser numérico, sem fórmula.")
    try:
        number = Decimal(str(value))
        if not number.is_finite():
            raise InvalidOperation
        return number.quantize(CENT, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError) as exc:
        raise ReportError(f"{cell}: valor monetário inválido.") from exc


def read_reports(path: Path | BinaryIO, settings: Settings, card: str, month: str) -> list[Report]:
    """Read primitive source cells; never evaluate formulas or trust their caches."""
    try:
        workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    except Exception as exc:
        raise ReportError("Não foi possível abrir o Excel. Verifique se o arquivo está íntegro e a cópia está completa.") from exc
    try:
        expected = {"Black", "Latam"} | {r.sheet for r in settings.recipients}
        unknown = set(workbook.sheetnames) - expected
        if unknown:
            raise ReportError("Abas sem cadastro: " + ", ".join(sorted(unknown)))
        missing = expected - set(workbook.sheetnames)
        if missing:
            raise ReportError("Abas ausentes: " + ", ".join(sorted(missing)))
        source = workbook[CARDS[card][0]]
        # Materialize once: read-only random access reparses XML repeatedly.
        try:
            rows = list(source.iter_rows())
        except Exception as exc:
            # In read-only mode openpyxl parses cell values lazily. Keep raw
            # parser errors (which can contain cell contents) out of the logs.
            raise ReportError("Não foi possível ler as células do Excel. Verifique se o arquivo está íntegro e a cópia está completa.") from exc
        if len(rows) < 4 or len(rows[1]) < 6:
            raise ReportError(f"{source.title}: layout inválido; cabeçalhos esperados na linha 2 a partir de F.")
        headers: dict[str, int] = {}
        for i, cell in enumerate(rows[1][5:], start=5):
            if cell.value is None:
                continue
            key = str(cell.value).strip().casefold()
            if key in headers:
                raise ReportError(f"Cabeçalho duplicado em {source.title}, linha 2.")
            headers[key] = i
        labels = [(str(row[0].value or "").strip().upper(), i) for i, row in enumerate(rows)]
        totals = [i for label, i in labels if label == "TOTAL"]
        balances = [i for label, i in labels if label == "TOTAL GERAL"]
        if len(totals) != 1 or len(balances) != 1 or totals[0] >= balances[0]:
            raise ReportError(f"{source.title}: esperava uma linha TOTAL seguida de TOTAL GERAL.")
        reports = []
        for recipient in settings.recipients:
            if not recipient.enabled:
                continue
            column = headers.get(recipient.participant.casefold())
            if column is None:
                raise ReportError(f"Coluna de rateio não encontrada para {recipient.sheet} em {source.title}.")
            purchases, payments = [], []
            for i, row in enumerate(rows[3:], start=3):
                label = str(row[0].value or "").strip().upper()
                if label in {"TOTAL", "TOTAL GERAL"}:
                    continue
                value = row[column].value
                if value is None:
                    continue
                cell_label = f"{source.title}!{openpyxl.utils.get_column_letter(column + 1)}{i + 1}"
                share = _money(value, cell_label)
                if share == ZERO:
                    continue
                if label == "PAGAMENTOS":
                    if not totals[0] < i < balances[0]:
                        raise ReportError(f"{cell_label}: pagamento fora do bloco esperado.")
                    payments.append(share)
                    continue
                when = row[0].value
                if not isinstance(when, date) or i >= totals[0]:
                    raise ReportError(f"{cell_label}: rateio sem data de compra válida ou fora do bloco de compras.")
                description = row[1].value
                if not isinstance(description, str) or not description.strip() or row[1].data_type in {"f", "e"}:
                    raise ReportError(f"{source.title}!B{i + 1}: lançamento inválido.")
                installment = row[2].value
                if row[2].data_type in {"f", "e"}:
                    raise ReportError(f"{source.title}!C{i + 1}: parcela contém fórmula ou erro.")
                purchases.append(Purchase(
                    when.date() if isinstance(when, datetime) else when,
                    description.strip(), "—" if installment is None or installment == 0 else str(installment),
                    _money(row[3].value, f"{source.title}!D{i + 1}"), share,
                ))
            reports.append(Report(card, month, recipient, tuple(purchases), tuple(payments), settings.owner_name, settings.payment_footer))
        return reports
    finally:
        workbook.close()
